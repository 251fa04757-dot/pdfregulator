"""
FastAPI Server for VFSTR Academic Document Formatting & Modifier Agent
Supports single and batch (multiple) Word document processing.
Runs locally on http://localhost:8000
"""

import os
import shutil
import zipfile
from pathlib import Path
from typing import Optional, List, Dict, Any
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse

import re
import json
import docx
import pymupdf

from academic_formatter import convert_docx_to_pdf
from academic_modifier import AcademicModifierAgent

BASE_DIR = Path(__file__).resolve().parent
UPLOADS_DIR = BASE_DIR / "uploads"
OUTPUTS_DIR = BASE_DIR / "outputs"
TEMPLATES_DIR = BASE_DIR / "templates"
ASSETS_DIR = BASE_DIR / "assets"

UPLOADS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ASSETS_DIR.mkdir(exist_ok=True)

app = FastAPI(title="VFSTR Academic Document Formatting & Modifier Agent")

# Memory store for active documents and their modifier agents
ACTIVE_STATE: Dict[str, Any] = {
    "active_docx": None,
    "active_pdf": None,
    "agents": {},  # docx_name -> AcademicModifierAgent
    "batch_documents": []
}


def build_agent_zip() -> Path:
    """Creates a standalone zip archive of the agent for offline distribution."""
    zip_path = BASE_DIR / "VFSTR_Academic_Document_Agent.zip"

    files_to_pack = [
        "academic_formatter.py",
        "academic_modifier.py",
        "app.py",
        "start_agent.bat",
        "requirements.txt",
        "README.md",
    ]

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for fname in files_to_pack:
            fpath = BASE_DIR / fname
            if fpath.exists():
                zf.write(fpath, arcname=fname)

        if TEMPLATES_DIR.exists():
            for t_file in TEMPLATES_DIR.glob("*"):
                zf.write(t_file, arcname=f"templates/{t_file.name}")

        if ASSETS_DIR.exists():
            for a_file in ASSETS_DIR.glob("*"):
                zf.write(a_file, arcname=f"assets/{a_file.name}")

    return zip_path


def build_batch_zip(pdf_paths: List[Path]) -> Path:
    """Creates a zip archive containing all formatted PDFs from a batch run."""
    batch_zip_path = OUTPUTS_DIR / "VFSTR_Batch_Formatted_PDFs.zip"
    temp_zip = OUTPUTS_DIR / "VFSTR_Batch_Formatted_PDFs.tmp.zip"
    try:
        with zipfile.ZipFile(temp_zip, "w", zipfile.ZIP_DEFLATED) as zf:
            for pdf_p in pdf_paths:
                if pdf_p.exists():
                    zf.write(pdf_p, arcname=pdf_p.name)
        if batch_zip_path.exists():
            try: batch_zip_path.unlink()
            except Exception: pass
        try:
            temp_zip.replace(batch_zip_path)
        except Exception:
            shutil.copy2(temp_zip, batch_zip_path)
            try: temp_zip.unlink()
            except Exception: pass
    except Exception as e:
        print(f"[build_batch_zip] Warning: {e}")
    return batch_zip_path


def merge_curriculum_pdfs(pdf_paths: List[Path], output_path: Path) -> Path:
    """Merges multiple individual curriculum PDFs sequentially so every subject starts on a fresh page."""
    combined_doc = pymupdf.open()
    for p in pdf_paths:
        if p.exists():
            src_doc = pymupdf.open(str(p))
            combined_doc.insert_pdf(src_doc)
            src_doc.close()
    
    temp_out = output_path.with_name(f"{output_path.stem}.tmp.pdf")
    try:
        combined_doc.save(str(temp_out))
        combined_doc.close()
        if output_path.exists():
            try: output_path.unlink()
            except Exception: pass
        try:
            temp_out.replace(output_path)
        except Exception:
            shutil.copy2(temp_out, output_path)
            try: temp_out.unlink()
            except Exception: pass
    except Exception as e:
        print(f"[merge_curriculum_pdfs] Warning: {e}")
        try: combined_doc.close()
        except Exception: pass
    return output_path


@app.get("/", response_class=FileResponse)
async def index():
    """Serve the main web UI."""
    html_file = TEMPLATES_DIR / "index.html"
    return FileResponse(html_file, media_type="text/html")


def convert_pdf_to_docx(pdf_path: str, output_docx_path: str):
    """Extracts text lines from a PDF and creates a clean docx document."""
    doc_pdf = pymupdf.open(pdf_path)
    doc_word = docx.Document()
    for page in doc_pdf:
        text = page.get_text()
        for line in text.split("\n"):
            line_str = line.strip()
            if line_str:
                doc_word.add_paragraph(line_str)
    doc_pdf.close()
    doc_word.save(output_docx_path)


@app.post("/convert")
async def convert_documents(
    files: List[UploadFile] = File(...),
    start_page: int = Form(80),
    continuous_pages: bool = Form(True),
    header_text: str = Form("CSE - II Year II Semester"),
    course_code: Optional[str] = Form(None),
    image_query: Optional[str] = Form(None),
    unsplash_api_key: Optional[str] = Form(None),
    skills: Optional[str] = Form(None)
):
    """Converts one or multiple uploaded Word documents (.docx) or PDFs (.pdf) to VFSTR formatted PDFs."""
    if not files:
        return JSONResponse(status_code=400, content={"success": False, "error": "No files provided."})

    valid_files = [f for f in files if f.filename and (f.filename.lower().endswith(".docx") or f.filename.lower().endswith(".pdf"))]
    if not valid_files:
        return JSONResponse(status_code=400, content={"success": False, "error": "Only Microsoft Word (.docx) and PDF (.pdf) files are supported."})

    skills_list = [s.strip() for s in skills.split("\n") if s.strip()] if skills else None
    results = []
    generated_pdf_paths = []
    current_page = start_page

    try:
        for f in valid_files:
            safe_name = Path(f.filename).stem
            input_ext = Path(f.filename).suffix.lower()
            input_docx_path = UPLOADS_DIR / f"{safe_name}.docx"
            output_pdf_filename = f"{safe_name}_VFSTR_Formatted.pdf"
            output_pdf_path = OUTPUTS_DIR / output_pdf_filename

            # Save uploaded file
            if input_ext == ".pdf":
                temp_pdf = UPLOADS_DIR / f"{safe_name}_raw.pdf"
                with open(temp_pdf, "wb") as buffer:
                    shutil.copyfileobj(f.file, buffer)
                convert_pdf_to_docx(str(temp_pdf), str(input_docx_path))
                if temp_pdf.exists():
                    try: os.remove(temp_pdf)
                    except Exception: pass
            else:
                with open(input_docx_path, "wb") as buffer:
                    shutil.copyfileobj(f.file, buffer)

            # Convert to VFSTR PDF
            convert_docx_to_pdf(
                docx_path=str(input_docx_path),
                output_pdf_path=str(output_pdf_path),
                start_page=current_page,
                header_text=header_text,
                course_code=course_code,
                subject_image_query=image_query,
                unsplash_api_key=unsplash_api_key,
                custom_skills=skills_list
            )

            # Determine number of pages
            page_count = 1
            try:
                doc = pymupdf.open(str(output_pdf_path))
                page_count = len(doc)
            except Exception:
                pass

            end_page = current_page + page_count - 1

            # Initialize Modifier Agent for this document
            agent = AcademicModifierAgent(
                docx_path=str(input_docx_path),
                output_pdf_path=str(output_pdf_path)
            )
            agent.current_settings["start_page"] = current_page
            agent.current_settings["header_text"] = header_text
            if course_code:
                agent.current_settings["course_code"] = course_code
            if image_query:
                agent.current_settings["subject_image_query"] = image_query
            if unsplash_api_key:
                agent.current_settings["unsplash_api_key"] = unsplash_api_key
            if skills_list:
                agent.current_settings["skills"] = skills_list

            ACTIVE_STATE["agents"][f"{safe_name}.docx"] = agent

            doc_info = {
                "title": safe_name.replace("_", " "),
                "docx_name": f"{safe_name}.docx",
                "pdf_filename": output_pdf_filename,
                "pdf_url": f"/view-pdf/{output_pdf_filename}",
                "download_url": f"/download-pdf/{output_pdf_filename}",
                "start_page": current_page,
                "end_page": end_page,
                "page_count": page_count,
                "active_settings": agent.current_settings
            }
            results.append(doc_info)
            generated_pdf_paths.append(output_pdf_path)

            # Advance page number for next document if continuous pagination requested
            if continuous_pages:
                current_page = end_page + 1

        ACTIVE_STATE["batch_documents"] = results
        ACTIVE_STATE["active_docx"] = results[0]["docx_name"]
        ACTIVE_STATE["active_pdf"] = results[0]["pdf_filename"]

        batch_zip_url = None
        combined_pdf_url = None
        combined_download_url = None
        combined_filename = None

        if len(generated_pdf_paths) > 1:
            combined_path = OUTPUTS_DIR / "Combined_Curriculum_VFSTR_Formatted.pdf"
            merge_curriculum_pdfs(generated_pdf_paths, combined_path)
            build_batch_zip(generated_pdf_paths + [combined_path])
            batch_zip_url = "/download-batch-zip"
            combined_filename = "Combined_Curriculum_VFSTR_Formatted.pdf"
            combined_pdf_url = "/view-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"
            combined_download_url = "/download-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"

        return {
            "success": True,
            "total_files": len(results),
            "documents": results,
            "batch_zip_url": batch_zip_url,
            "combined_pdf_url": combined_pdf_url,
            "combined_download_url": combined_download_url,
            "combined_filename": combined_filename,
            "active_doc": results[0]
        }

    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


def parse_docx_to_dict(docx_path: str) -> dict:
    """Parses a Word document into structured fields for editing."""
    doc = docx.Document(docx_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    raw_text = "\n".join(paras)

    hours = {'L': '0', 'T': '0', 'P': '0', 'SL': '0', 'C': '3'}
    for t in doc.tables:
        if len(t.rows) >= 2 and len(t.columns) == 5:
            hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
            if any(h in ['L', 'T', 'P', 'C'] for h in hdrs):
                for h, cell in zip(hdrs, t.rows[1].cells):
                    hours[h] = cell.text.strip()
                break

    title = paras[0] if paras else ''
    code_m = re.search(r'\b([0-9]{2}[A-Z]{2,4}[0-9]{3,4}[A-Z]?)\b', title)
    code = code_m.group(1).upper() if code_m else ''
    if code:
        title = title.replace(code, '').strip(' :-_\t')

    prereq = ''
    desc_lines = []
    mod_lines = []
    tb_lines = []
    rb_lines = []
    skills_lines = []

    section = 'header'
    for p in paras[1:]:
        up = p.upper()
        if up.startswith('PREREQUISITE KNOWLEDGE:'):
            prereq = p.split(':', 1)[-1].strip()
            continue
        elif 'COURSE DESCRIPTION AND OBJECTIVES' in up:
            section = 'desc'
            continue
        elif re.match(r'^MODULE\s*[-–—]?\s*[0-9IVX]+', up):
            section = 'modules'
        elif 'TEXT BOOKS' in up:
            section = 'textbooks'
            continue
        elif 'REFERENCE BOOKS' in up:
            section = 'refbooks'
            continue
        elif up.startswith('SKILLS:'):
            section = 'skills'
            continue
        elif 'COURSE OUTCOMES' in up:
            section = 'outcomes'
            continue

        if section == 'desc':
            desc_lines.append(p)
        elif section == 'modules':
            mod_lines.append(p)
        elif section == 'textbooks':
            tb_lines.append(p)
        elif section == 'refbooks':
            rb_lines.append(p)
        elif section == 'skills':
            clean_s = re.sub(r'^[✓✔\-\*•●]\s*', '', p).strip()
            if clean_s:
                skills_lines.append(clean_s)

    course_outcomes_table = []
    sdgs_table = []

    for t in doc.tables:
        if len(t.rows) >= 2 and len(t.columns) == 5:
            hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
            if not any(h in ['L', 'T', 'P', 'SL', 'C'] for h in hdrs):
                for row in t.rows:
                    course_outcomes_table.append([c.text.strip() for c in row.cells])
        elif len(t.rows) >= 2:
            hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
            if any('SDG' in h or 'IKS' in h or 'SUSTAINABLE' in h for h in hdrs):
                for row in t.rows:
                    sdgs_table.append([c.text.strip() for c in row.cells])

    if not course_outcomes_table:
        for t in doc.tables:
            if len(t.rows) >= 2:
                hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
                if any('OUTCOME' in h for h in hdrs):
                    for row in t.rows:
                        course_outcomes_table.append([c.text.strip() for c in row.cells])
                    break

    sdg_numbers = set()
    for row in sdgs_table:
        for cell in row:
            for m in re.findall(r'(?:SDG|Goal)\s*[-:]?\s*([0-9]{1,2})', cell, re.I):
                try:
                    num = int(m)
                    if 1 <= num <= 17:
                        sdg_numbers.add(num)
                except ValueError:
                    pass
    if not sdg_numbers:
        for m in re.findall(r'(?:SDG|Goal)\s*[-:]?\s*([0-9]{1,2})', raw_text, re.I):
            try:
                num = int(m)
                if 1 <= num <= 17:
                    sdg_numbers.add(num)
            except ValueError:
                pass

    return {
        "course_title": title,
        "course_code": code,
        "hours": hours,
        "prerequisite": prereq,
        "description": "\n\n".join(desc_lines),
        "modules_content": "\n".join(mod_lines),
        "textbooks": "\n".join(tb_lines),
        "references": "\n".join(rb_lines),
        "skills": skills_lines,
        "course_outcomes_table": course_outcomes_table,
        "sdgs_table": sdgs_table,
        "sdg_numbers": sorted(list(sdg_numbers)),
        "raw_text": raw_text
    }


def update_docx_from_dict(docx_path: str, data: dict):
    """Updates paragraphs and tables in docx with edited content."""
    doc = docx.Document(docx_path)

    # 1. Title and Course Code
    course_title = data.get("course_title")
    course_code = data.get("course_code")
    if course_title:
        title_text = course_title.strip()
        if course_code:
            title_text = f"{course_code.strip()} {title_text}"
        for p in doc.paragraphs:
            if p.text.strip():
                p.text = title_text
                break

    # 2. Hours Table
    hours = data.get("hours", {})
    if hours:
        for t in doc.tables:
            if len(t.rows) >= 2 and len(t.columns) == 5:
                hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
                if any(h in ['L', 'T', 'P', 'C'] for h in hdrs):
                    for idx, h in enumerate(hdrs):
                        if h in hours and hours[h] is not None:
                            t.rows[1].cells[idx].text = str(hours[h])
                    break

    # 3. Prerequisite
    prerequisite = data.get("prerequisite")
    if prerequisite is not None:
        found_prereq = False
        for p in doc.paragraphs:
            if p.text.strip().upper().startswith("PREREQUISITE KNOWLEDGE:"):
                p.text = f"PREREQUISITE KNOWLEDGE: {prerequisite.strip()}"
                found_prereq = True
                break
        if not found_prereq and prerequisite.strip():
            doc.add_paragraph(f"PREREQUISITE KNOWLEDGE: {prerequisite.strip()}")

    # 4. Description
    description = data.get("description")
    if description is not None:
        desc_found = False
        for i, p in enumerate(doc.paragraphs):
            if "COURSE DESCRIPTION AND OBJECTIVES" in p.text.strip().upper():
                if i + 1 < len(doc.paragraphs):
                    doc.paragraphs[i + 1].text = description.strip()
                    desc_found = True
                    break
        if not desc_found and description.strip():
            doc.add_paragraph("COURSE DESCRIPTION AND OBJECTIVES:")
            doc.add_paragraph(description.strip())

    # 5. Skills
    skills = data.get("skills")
    if skills:
        skills_start_idx = None
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip().upper().startswith("SKILLS:"):
                skills_start_idx = i
                break
        if skills_start_idx is not None:
            curr = skills_start_idx + 1
            for s_item in skills:
                if curr < len(doc.paragraphs):
                    doc.paragraphs[curr].text = f"✓ {s_item}"
                    curr += 1
                else:
                    doc.add_paragraph(f"✓ {s_item}")
        else:
            doc.add_paragraph("SKILLS:")
            for s_item in skills:
                doc.add_paragraph(f"✓ {s_item}")

    # 6. Textbooks
    textbooks = data.get("textbooks")
    if textbooks is not None and textbooks.strip():
        tb_idx = None
        for i, p in enumerate(doc.paragraphs):
            if "TEXTBOOK" in p.text.strip().upper() or "TEXT BOOK" in p.text.strip().upper():
                tb_idx = i
                break
        if tb_idx is not None and tb_idx + 1 < len(doc.paragraphs):
            doc.paragraphs[tb_idx + 1].text = textbooks.strip()

    # 7. Reference Books
    references = data.get("references")
    if references is not None and references.strip():
        rb_idx = None
        for i, p in enumerate(doc.paragraphs):
            if "REFERENCE BOOK" in p.text.strip().upper() or "REFERENCES:" in p.text.strip().upper():
                rb_idx = i
                break
        if rb_idx is not None and rb_idx + 1 < len(doc.paragraphs):
            doc.paragraphs[rb_idx + 1].text = references.strip()

    # 8. Course Outcomes Table (Inline editable table)
    course_outcomes_table = data.get("course_outcomes_table")
    if course_outcomes_table and len(course_outcomes_table) >= 2:
        co_table = None
        for t in doc.tables:
            if len(t.rows) >= 1 and len(t.columns) == 5:
                hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
                if not any(h in ['L', 'T', 'P', 'SL', 'C'] for h in hdrs):
                    co_table = t
                    break
        if co_table is None:
            for t in doc.tables:
                if len(t.rows) >= 1:
                    hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
                    if any('OUTCOME' in h for h in hdrs):
                        co_table = t
                        break

        if co_table:
            target_num_rows = len(course_outcomes_table)
            while len(co_table.rows) < target_num_rows:
                co_table.add_row()
            while len(co_table.rows) > target_num_rows and len(co_table.rows) > 2:
                r = co_table.rows[-1]
                r._tr.getparent().remove(r._tr)

            for r_idx, row_vals in enumerate(course_outcomes_table):
                if r_idx < len(co_table.rows):
                    row_obj = co_table.rows[r_idx]
                    for c_idx, val in enumerate(row_vals):
                        if c_idx < len(row_obj.cells):
                            row_obj.cells[c_idx].text = str(val or "")

    # 9. SDGs & IKS Table
    sdgs_table = data.get("sdgs_table")
    if sdgs_table and len(sdgs_table) >= 2:
        sdg_table = None
        for t in doc.tables:
            if len(t.rows) >= 1:
                hdrs = [c.text.strip().upper() for c in t.rows[0].cells]
                if any('SDG' in h or 'IKS' in h or 'SUSTAINABLE' in h for h in hdrs):
                    sdg_table = t
                    break
        if sdg_table:
            target_num_rows = len(sdgs_table)
            while len(sdg_table.rows) < target_num_rows:
                sdg_table.add_row()
            while len(sdg_table.rows) > target_num_rows and len(sdg_table.rows) > 2:
                r = sdg_table.rows[-1]
                r._tr.getparent().remove(r._tr)

            for r_idx, row_vals in enumerate(sdgs_table):
                if r_idx < len(sdg_table.rows):
                    row_obj = sdg_table.rows[r_idx]
                    for c_idx, val in enumerate(row_vals):
                        if c_idx < len(row_obj.cells):
                            row_obj.cells[c_idx].text = str(val or "")

    doc.save(docx_path)


@app.get("/get-document-content")
async def get_document_content(docx_name: Optional[str] = None):
    """Retrieves full parsed content and settings for editing."""
    target_docx = docx_name or ACTIVE_STATE.get("active_docx")
    if not target_docx:
        docx_files = sorted(list(UPLOADS_DIR.glob("*.docx")), key=os.path.getmtime, reverse=True)
        if docx_files:
            target_docx = docx_files[0].name
        else:
            return JSONResponse(status_code=400, content={"success": False, "error": "No document found."})

    docx_path = UPLOADS_DIR / target_docx
    if not docx_path.exists():
        return JSONResponse(status_code=404, content={"success": False, "error": f"Document {target_docx} not found."})

    parsed = parse_docx_to_dict(str(docx_path))
    agent = ACTIVE_STATE.get("agents", {}).get(target_docx)
    settings = agent.current_settings if agent else {}

    doc_info = next((d for d in ACTIVE_STATE.get("batch_documents", []) if d["docx_name"] == target_docx), None)

    return {
        "success": True,
        "docx_name": target_docx,
        "course_title": parsed["course_title"],
        "course_code": parsed["course_code"] or settings.get("course_code", "24CS206"),
        "header_text": settings.get("header_text", "CSE - II Year II Semester"),
        "start_page": settings.get("start_page", doc_info.get("start_page", 80) if doc_info else 80),
        "hours": parsed["hours"],
        "prerequisite": parsed["prerequisite"],
        "description": parsed["description"],
        "modules_content": parsed["modules_content"],
        "textbooks": parsed["textbooks"],
        "references": parsed["references"],
        "skills": parsed["skills"] or settings.get("skills", []),
        "course_outcomes_table": parsed.get("course_outcomes_table", []),
        "sdgs_table": parsed.get("sdgs_table", []),
        "custom_source_url": settings.get("custom_source_url", ""),
        "custom_image_url": settings.get("custom_image_url", ""),
        "subject_image_query": settings.get("subject_image_query", ""),
        "raw_text": parsed["raw_text"]
    }


@app.post("/save-pdf-content")
async def save_pdf_content(
    docx_name: str = Form(...),
    course_title: Optional[str] = Form(None),
    course_code: Optional[str] = Form(None),
    header_text: Optional[str] = Form(None),
    start_page: Optional[int] = Form(None),
    hours_l: Optional[str] = Form(None),
    hours_t: Optional[str] = Form(None),
    hours_p: Optional[str] = Form(None),
    hours_sl: Optional[str] = Form(None),
    hours_c: Optional[str] = Form(None),
    prerequisite: Optional[str] = Form(None),
    description: Optional[str] = Form(None),
    modules_content: Optional[str] = Form(None),
    textbooks: Optional[str] = Form(None),
    references: Optional[str] = Form(None),
    skills: Optional[str] = Form(None),
    course_outcomes_table: Optional[str] = Form(None),
    sdgs_table: Optional[str] = Form(None),
    custom_source_url: Optional[str] = Form(None),
    custom_image_url: Optional[str] = Form(None),
    subject_image_query: Optional[str] = Form(None),
    image_file: Optional[UploadFile] = File(None)
):
    """Saves all user-edited sections, rebuilds document, re-compiles PDF, and updates batch combined book."""
    target_docx = docx_name or ACTIVE_STATE.get("active_docx")
    if not target_docx:
        return JSONResponse(status_code=400, content={"success": False, "error": "No document specified."})

    docx_path = UPLOADS_DIR / target_docx
    safe_name = Path(target_docx).stem
    output_pdf_filename = f"{safe_name}_VFSTR_Formatted.pdf"
    output_pdf_path = OUTPUTS_DIR / output_pdf_filename

    hours = {}
    if hours_l is not None: hours["L"] = hours_l
    if hours_t is not None: hours["T"] = hours_t
    if hours_p is not None: hours["P"] = hours_p
    if hours_sl is not None: hours["SL"] = hours_sl
    if hours_c is not None: hours["C"] = hours_c

    skills_list = [s.strip() for s in skills.split("\n") if s.strip()] if skills else None

    # Parse course outcomes table JSON if provided
    co_table_data = None
    if course_outcomes_table:
        try:
            co_table_data = json.loads(course_outcomes_table)
        except Exception as e:
            print(f"[save-pdf-content] Error parsing course_outcomes_table: {e}")

    sdgs_table_data = None
    if sdgs_table:
        try:
            sdgs_table_data = json.loads(sdgs_table)
        except Exception as e:
            print(f"[save-pdf-content] Error parsing sdgs_table: {e}")

    # Update docx file on disk
    update_data = {
        "course_title": course_title,
        "course_code": course_code,
        "hours": hours,
        "prerequisite": prerequisite,
        "description": description,
        "textbooks": textbooks,
        "references": references,
        "skills": skills_list,
        "course_outcomes_table": co_table_data,
        "sdgs_table": sdgs_table_data
    }
    update_docx_from_dict(str(docx_path), update_data)

    agent = ACTIVE_STATE["agents"].get(target_docx)
    if not agent or agent.docx_path != str(docx_path):
        agent = AcademicModifierAgent(str(docx_path), str(output_pdf_path))
        ACTIVE_STATE["agents"][target_docx] = agent

    overrides = {}
    if start_page is not None:
        overrides["start_page"] = start_page
    if header_text:
        overrides["header_text"] = header_text
    if course_code:
        overrides["course_code"] = course_code
    if subject_image_query:
        overrides["subject_image_query"] = subject_image_query
    elif course_title:
        overrides["subject_image_query"] = course_title
    if custom_image_url:
        overrides["custom_image_url"] = custom_image_url
    if custom_source_url:
        overrides["custom_source_url"] = custom_source_url.strip()
    if skills_list:
        overrides["skills"] = skills_list

    # Handle image file upload from device
    if image_file and image_file.filename:
        custom_img_dir = UPLOADS_DIR / "custom_images"
        custom_img_dir.mkdir(parents=True, exist_ok=True)
        suffix = Path(image_file.filename).suffix or ".jpg"
        saved_img_path = custom_img_dir / f"{safe_name}_uploaded_img{suffix}"
        with open(saved_img_path, "wb") as f_img:
            shutil.copyfileobj(image_file.file, f_img)
        overrides["custom_image_path"] = str(saved_img_path)

    result = agent.modify(instruction="Save modified PDF with updated curriculum content", overrides=overrides)

    ACTIVE_STATE["active_docx"] = target_docx
    ACTIVE_STATE["active_pdf"] = output_pdf_filename

    # Recalculate page count
    page_count = 1
    try:
        doc = pymupdf.open(str(output_pdf_path))
        page_count = len(doc)
        doc.close()
    except Exception:
        pass

    for d in ACTIVE_STATE.get("batch_documents", []):
        if d["docx_name"] == target_docx:
            d["page_count"] = page_count
            d["end_page"] = d["start_page"] + page_count - 1
            d["active_settings"] = agent.current_settings

    all_pdfs = [OUTPUTS_DIR / d["pdf_filename"] for d in ACTIVE_STATE.get("batch_documents", []) if (OUTPUTS_DIR / d["pdf_filename"]).exists()]
    if len(all_pdfs) > 1:
        combined_path = OUTPUTS_DIR / "Combined_Curriculum_VFSTR_Formatted.pdf"
        merge_curriculum_pdfs(all_pdfs, combined_path)
        build_batch_zip(all_pdfs + [combined_path])
        result["combined_pdf_url"] = "/view-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"
        result["combined_download_url"] = "/download-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"
        result["combined_filename"] = "Combined_Curriculum_VFSTR_Formatted.pdf"

    result["message"] = f"Modified PDF for '{target_docx}' saved successfully!"
    result["pdf_url"] = f"/view-pdf/{output_pdf_filename}"
    result["download_url"] = f"/download-pdf/{output_pdf_filename}"
    result["filename"] = output_pdf_filename
    return result


@app.post("/modify")
async def modify_document(
    docx_name: Optional[str] = Form(None),
    instruction: Optional[str] = Form(""),
    start_page: Optional[int] = Form(None),
    header_text: Optional[str] = Form(None),
    course_code: Optional[str] = Form(None),
    image_query: Optional[str] = Form(None),
    custom_image_url: Optional[str] = Form(None),
    custom_source_url: Optional[str] = Form(None),
    unsplash_api_key: Optional[str] = Form(None),
    skills: Optional[str] = Form(None),
    image_file: Optional[UploadFile] = File(None)
):
    """Invokes the AI Modifier Agent to modify a specific document from the batch."""
    target_docx = docx_name or ACTIVE_STATE["active_docx"]

    if not target_docx:
        docx_files = sorted(list(UPLOADS_DIR.glob("*.docx")), key=os.path.getmtime, reverse=True)
        if docx_files:
            target_docx = docx_files[0].name
        else:
            return JSONResponse(status_code=400, content={"success": False, "error": "No Word document found to modify."})

    docx_path = UPLOADS_DIR / target_docx
    safe_name = Path(target_docx).stem
    output_pdf_filename = f"{safe_name}_VFSTR_Formatted.pdf"
    output_pdf_path = OUTPUTS_DIR / output_pdf_filename

    # Retrieve or create agent
    agent = ACTIVE_STATE["agents"].get(target_docx)
    if not agent or agent.docx_path != str(docx_path):
        agent = AcademicModifierAgent(
            docx_path=str(docx_path),
            output_pdf_path=str(output_pdf_path)
        )
        ACTIVE_STATE["agents"][target_docx] = agent

    overrides = {}
    if start_page is not None:
        overrides["start_page"] = start_page
    if header_text:
        overrides["header_text"] = header_text
    if course_code:
        overrides["course_code"] = course_code
    if image_query:
        overrides["subject_image_query"] = image_query
    if custom_image_url:
        overrides["custom_image_url"] = custom_image_url
    if custom_source_url and custom_source_url.strip():
        overrides["custom_source_url"] = custom_source_url.strip()
    if unsplash_api_key:
        overrides["unsplash_api_key"] = unsplash_api_key
    if skills:
        items = [s.strip() for s in skills.split("\n") if s.strip()]
        overrides["skills"] = items

    # Handle image file uploaded from local computer folder
    if image_file and image_file.filename:
        custom_img_dir = UPLOADS_DIR / "custom_images"
        custom_img_dir.mkdir(parents=True, exist_ok=True)
        suffix = Path(image_file.filename).suffix or ".jpg"
        saved_img_path = custom_img_dir / f"{safe_name}_uploaded_img{suffix}"
        with open(saved_img_path, "wb") as f_img:
            shutil.copyfileobj(image_file.file, f_img)
        overrides["custom_image_path"] = str(saved_img_path)

    try:
        result = agent.modify(instruction=instruction, overrides=overrides)
        ACTIVE_STATE["active_docx"] = target_docx
        ACTIVE_STATE["active_pdf"] = output_pdf_filename

        # If batch mode, update the batch doc entry and re-merge combined book & zip
        if len(ACTIVE_STATE.get("batch_documents", [])) > 1:
            for d in ACTIVE_STATE["batch_documents"]:
                if d["docx_name"] == target_docx:
                    d["active_settings"] = agent.current_settings

            all_pdfs = [OUTPUTS_DIR / d["pdf_filename"] for d in ACTIVE_STATE["batch_documents"] if (OUTPUTS_DIR / d["pdf_filename"]).exists()]
            if len(all_pdfs) > 1:
                combined_path = OUTPUTS_DIR / "Combined_Curriculum_VFSTR_Formatted.pdf"
                merge_curriculum_pdfs(all_pdfs, combined_path)
                build_batch_zip(all_pdfs + [combined_path])
                result["combined_pdf_url"] = "/view-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"
                result["combined_download_url"] = "/download-pdf/Combined_Curriculum_VFSTR_Formatted.pdf"
                result["combined_filename"] = "Combined_Curriculum_VFSTR_Formatted.pdf"

        return result
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@app.post("/reorder-batch")
async def reorder_batch(request: Request):
    """Reorders the batch documents according to the user's custom sequence,
    re-paginates continuously, and regenerates the combined curriculum book."""
    try:
        content_type = request.headers.get("content-type", "")
        if "application/json" in content_type:
            body = await request.json()
            order = body.get("order", [])
            start_page = int(body.get("start_page", 80))
            continuous_pages = bool(body.get("continuous_pages", True))
        else:
            form = await request.form()
            order = form.get("order", "")
            start_page = int(form.get("start_page", 80))
            continuous_pages = str(form.get("continuous_pages", "true")).lower() in ("true", "1")

        if isinstance(order, list):
            ordered_names = [str(x).strip() for x in order if str(x).strip()]
        elif isinstance(order, str):
            if order.strip().startswith("["):
                ordered_names = json.loads(order)
            else:
                ordered_names = [s.strip() for s in order.split(",") if s.strip()]
        else:
            ordered_names = []

        batch_docs = ACTIVE_STATE.get("batch_documents", [])
        if not batch_docs:
            return JSONResponse(status_code=400, content={"success": False, "error": "No batch documents to reorder."})

        docs_map = {d["docx_name"]: d for d in batch_docs}
        new_batch_docs = []
        for name in ordered_names:
            if name in docs_map and docs_map[name] not in new_batch_docs:
                new_batch_docs.append(docs_map[name])
        for d in batch_docs:
            if d not in new_batch_docs:
                new_batch_docs.append(d)

        current_page = start_page
        generated_pdf_paths = []

        for d in new_batch_docs:
            name = d["docx_name"]
            agent = ACTIVE_STATE["agents"].get(name)
            docx_path = UPLOADS_DIR / name
            output_pdf_path = OUTPUTS_DIR / d["pdf_filename"]

            if not agent:
                agent = AcademicModifierAgent(str(docx_path), str(output_pdf_path))
                ACTIVE_STATE["agents"][name] = agent

            # Re-compile with updated starting page number
            agent.modify(overrides={"start_page": current_page})

            # Check new page count
            page_count = d.get("page_count", 1)
            try:
                doc = pymupdf.open(str(output_pdf_path))
                page_count = len(doc)
                doc.close()
            except Exception:
                pass

            end_page = current_page + page_count - 1
            d["start_page"] = current_page
            d["end_page"] = end_page
            d["page_count"] = page_count
            d["active_settings"] = agent.current_settings

            generated_pdf_paths.append(output_pdf_path)

            if continuous_pages:
                current_page = end_page + 1

        ACTIVE_STATE["batch_documents"] = new_batch_docs
        ACTIVE_STATE["active_docx"] = new_batch_docs[0]["docx_name"]
        ACTIVE_STATE["active_pdf"] = new_batch_docs[0]["pdf_filename"]

        combined_path = OUTPUTS_DIR / "Combined_Curriculum_VFSTR_Formatted.pdf"
        merge_curriculum_pdfs(generated_pdf_paths, combined_path)
        build_batch_zip(generated_pdf_paths + [combined_path])

        return {
            "success": True,
            "documents": new_batch_docs,
            "combined_pdf_url": "/view-pdf/Combined_Curriculum_VFSTR_Formatted.pdf",
            "combined_download_url": "/download-pdf/Combined_Curriculum_VFSTR_Formatted.pdf",
            "active_doc": new_batch_docs[0]
        }
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@app.get("/view-pdf/{filename}")
async def view_pdf(filename: str):
    """Stream PDF for inline iframe preview."""
    file_path = OUTPUTS_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(file_path, media_type="application/pdf")


@app.get("/download-pdf/{filename}")
async def download_pdf(filename: str):
    """Download the generated PDF as an attachment."""
    file_path = OUTPUTS_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(file_path, media_type="application/pdf", filename=filename)


@app.get("/download-batch-zip")
async def download_batch_zip():
    """Download all PDFs generated from a multi-file batch run."""
    zip_path = OUTPUTS_DIR / "VFSTR_Batch_Formatted_PDFs.zip"
    if not zip_path.exists():
        raise HTTPException(status_code=404, detail="Batch zip not found")
    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename="VFSTR_Batch_Formatted_PDFs.zip"
    )


@app.get("/download-agent")
async def download_agent():
    """Download the standalone agent package (.zip)."""
    zip_path = build_agent_zip()
    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename="VFSTR_Academic_Document_Agent.zip"
    )


if __name__ == "__main__":
    import uvicorn
    print("Starting VFSTR Academic Document Formatting & Modifier Agent on http://localhost:8000 ...")
    uvicorn.run(app, host="127.0.0.1", port=8000)
